#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
登記所備付地図データ(GeoJSON変換済み)から、指定坪数レンジの土地(筆)を抽出してExcel化するツール。

使い方:
    python3 extract_parcels.py 入力1.geojson [入力2.geojson ...] \
        --min-tsubo 150 --max-tsubo 200 -o 浜松_150-200坪.xlsx

入力データの入手:
    G空間情報センター「法務省登記所備付地図データ」(無料・要アカウント登録)
    https://front.geospatial.jp/moj-chizu-xml-readme/
    → 変換済みデータ(GeoJSON)をダウンロードするか、地図XMLを
      mojxml2geojson (https://github.com/digital-go-jp/mojxml2geojson) で変換する。

注意:
    - 地図XMLには登記地積(公簿面積)は含まれない。本ツールはポリゴンの図上面積で
      近似するため、公簿面積とズレることがある。レンジは少し広めに取ること
      (例: 150〜200坪が欲しければ --buffer 0.1 で135〜220坪を拾う)。
    - 「任意座標系」の筆は正しい緯度経度を持たないため自動でスキップされる
      (地方の未整備地域に多い)。スキップ件数はログに出る。
"""

import argparse
import json
import math
import re
import sys
import zipfile
from pathlib import Path

TSUBO_PER_M2 = 1 / 3.305785  # 1坪 = 3.305785㎡

# 地番として無効・営業対象外にしやすいもの(道路・水路・筆界未定など)
SKIP_CHIBAN_PAT = re.compile(r"(道|水|畦|溝|堤|別図|地区外|未定)")

# 日本の緯度経度としてあり得る範囲(任意座標系データの混入ガード)
LON_RANGE = (122.0, 154.0)
LAT_RANGE = (20.0, 46.0)


def ring_area_m2(ring, lat0):
    """経緯度リングの面積を㎡で近似計算(局所平面近似+靴ひも公式)。

    ~1km四方以下の筆ポリゴンなら誤差0.1%未満で、坪レンジ抽出には十分。
    """
    m_per_deg_lat = 111132.954 - 559.822 * math.cos(2 * math.radians(lat0)) \
        + 1.175 * math.cos(4 * math.radians(lat0))
    m_per_deg_lon = 111412.84 * math.cos(math.radians(lat0)) \
        - 93.5 * math.cos(3 * math.radians(lat0))
    lon0, lat_base = ring[0][0], ring[0][1]
    area2 = 0.0
    n = len(ring)
    for i in range(n):
        x1 = (ring[i][0] - lon0) * m_per_deg_lon
        y1 = (ring[i][1] - lat_base) * m_per_deg_lat
        j = (i + 1) % n
        x2 = (ring[j][0] - lon0) * m_per_deg_lon
        y2 = (ring[j][1] - lat_base) * m_per_deg_lat
        area2 += x1 * y2 - x2 * y1
    return abs(area2) / 2.0


def polygon_area_m2(geom):
    """Polygon / MultiPolygon の面積(外周−穴)を㎡で返す。"""
    if geom is None:
        return None
    gtype = geom.get("type")
    if gtype == "Polygon":
        polys = [geom["coordinates"]]
    elif gtype == "MultiPolygon":
        polys = geom["coordinates"]
    else:
        return None
    total = 0.0
    for poly in polys:
        if not poly or not poly[0]:
            continue
        lat0 = poly[0][0][1]
        total += ring_area_m2(poly[0], lat0)          # 外周
        for hole in poly[1:]:
            total -= ring_area_m2(hole, lat0)         # 穴
    return total


def centroid(geom):
    """代表点(外周座標の単純平均)。地図リンク用途なので厳密な重心でなくてよい。"""
    gtype = geom.get("type")
    ring = geom["coordinates"][0] if gtype == "Polygon" else geom["coordinates"][0][0]
    lon = sum(p[0] for p in ring) / len(ring)
    lat = sum(p[1] for p in ring) / len(ring)
    return lat, lon


def prop(props, *names):
    """属性名のゆらぎ(変換ツールの版差)を吸収して値を取る。"""
    for n in names:
        v = props.get(n)
        if v not in (None, "", "null"):
            return str(v)
    return ""


def iter_features(path: Path):
    """GeoJSONファイル(.geojson/.json/.zip)からfeatureを順に返す。"""
    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if name.lower().endswith((".geojson", ".json")):
                    with zf.open(name) as f:
                        data = json.load(f)
                    yield from data.get("features", [])
    else:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        yield from data.get("features", [])


def main():
    ap = argparse.ArgumentParser(description="登記所備付地図GeoJSONから坪数レンジで筆を抽出しExcel化")
    ap.add_argument("inputs", nargs="+", help="GeoJSONファイル(.geojson/.json/.zip)")
    ap.add_argument("--min-tsubo", type=float, default=150.0)
    ap.add_argument("--max-tsubo", type=float, default=200.0)
    ap.add_argument("--buffer", type=float, default=0.0,
                    help="図上面積と公簿面積の誤差を見込んだレンジ拡張率(例: 0.1で±10%%広げる)")
    ap.add_argument("-o", "--output", default="parcels.xlsx", help="出力Excelファイル名")
    ap.add_argument("--csv", default=None, help="登記情報一括請求用CSVも出力する場合のファイル名")
    args = ap.parse_args()

    lo = args.min_tsubo * (1 - args.buffer)
    hi = args.max_tsubo * (1 + args.buffer)
    lo_m2 = lo / TSUBO_PER_M2
    hi_m2 = hi / TSUBO_PER_M2
    print(f"抽出レンジ: {lo:.1f}〜{hi:.1f}坪 ({lo_m2:.1f}〜{hi_m2:.1f}㎡)")

    rows = []
    n_total = n_arbitrary = n_skipped_chiban = 0

    for inp in args.inputs:
        path = Path(inp)
        if not path.exists():
            print(f"[警告] ファイルが見つかりません: {inp}", file=sys.stderr)
            continue
        print(f"読み込み中: {path.name}")
        for feat in iter_features(path):
            n_total += 1
            props = feat.get("properties") or {}
            geom = feat.get("geometry")
            chiban = prop(props, "地番", "chiban", "CHIBAN")
            if not chiban or SKIP_CHIBAN_PAT.search(chiban):
                n_skipped_chiban += 1
                continue
            area = polygon_area_m2(geom)
            if area is None or area <= 0:
                continue
            lat, lon = centroid(geom)
            # 任意座標系(ローカル座標)の筆を除外
            if not (LON_RANGE[0] <= lon <= LON_RANGE[1] and LAT_RANGE[0] <= lat <= LAT_RANGE[1]):
                n_arbitrary += 1
                continue
            if not (lo_m2 <= area <= hi_m2):
                continue
            city = prop(props, "市区町村名", "city_name", "CITY_NAME")
            oaza = prop(props, "大字名", "oaza_name", "OAZA_NAME")
            chome = prop(props, "丁目名", "chome_name", "CHOME_NAME")
            koaza = prop(props, "小字名", "koaza_name", "KOAZA_NAME")
            shozai = f"{city}{oaza}{chome}{koaza}"
            rows.append({
                "所在": shozai,
                "地番": chiban,
                "図上面積㎡": round(area, 1),
                "図上坪数": round(area * TSUBO_PER_M2, 1),
                "精度区分": prop(props, "精度区分", "seido_kubun"),
                "座標値種別": prop(props, "座標値種別", "zahyochi_shubetsu"),
                "緯度": round(lat, 7),
                "経度": round(lon, 7),
            })

    print(f"総筆数: {n_total} / 地番スキップ(道・水路等): {n_skipped_chiban} "
          f"/ 任意座標スキップ: {n_arbitrary} / 抽出: {len(rows)}")

    if not rows:
        print("該当なし。--buffer でレンジを広げるか、入力データを確認してください。")
        return

    rows.sort(key=lambda r: (r["所在"], r["地番"]))
    write_excel(rows, args.output)
    print(f"Excel出力: {args.output} ({len(rows)}件)")

    if args.csv:
        import csv as csvmod
        with open(args.csv, "w", newline="", encoding="cp932", errors="replace") as f:
            w = csvmod.writer(f)
            w.writerow(["所在", "地番"])
            for r in rows:
                w.writerow([r["所在"], r["地番"]])
        print(f"一括請求用CSV出力: {args.csv}")


def write_excel(rows, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "抽出リスト"
    headers = ["No", "所在", "地番", "図上面積㎡", "図上坪数", "精度区分",
               "座標値種別", "緯度", "経度", "地図リンク",
               "所有者(謄本取得後に記入)", "所有者住所", "地目", "備考"]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="1F4E79")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
    for i, r in enumerate(rows, start=1):
        link = f"https://www.google.com/maps?q={r['緯度']},{r['経度']}"
        ws.append([i, r["所在"], r["地番"], r["図上面積㎡"], r["図上坪数"],
                   r["精度区分"], r["座標値種別"], r["緯度"], r["経度"],
                   link, "", "", "", ""])
        cell = ws.cell(row=i + 1, column=10)
        cell.hyperlink = link
        cell.value = "地図で見る"
        cell.font = Font(color="0563C1", underline="single")
    widths = [5, 30, 10, 11, 10, 9, 12, 11, 11, 12, 22, 30, 8, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:N{len(rows) + 1}"
    wb.save(out_path)


if __name__ == "__main__":
    main()
