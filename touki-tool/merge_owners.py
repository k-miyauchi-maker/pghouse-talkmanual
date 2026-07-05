#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
所有者情報の一括取得結果(CSV)を、extract_parcels.py が出力したExcelに統合し、
DM・訪問営業用の所有者別リストまで一気に作るツール。全体フローの③を担当する。

    ① extract_parcels.py → 浜松_150-200坪.xlsx + 一括請求用.csv
    ② 一括取得サービス(ホームズ/登記簿図書館等) → 所有者CSV
    ③ このツール → 所有者入りExcel + 所有者別DMリスト

使い方:
    python3 merge_owners.py 浜松_150-200坪.xlsx 所有者.csv -o 浜松_DMリスト.xlsx

所有者CSVの形式:
    ヘッダー行に「所在」「地番」「所有者」(または氏名/名義人)、「所有者住所」(または住所)
    を含むCSV(UTF-8/Shift_JIS どちらでも可)。サービスの出力列名の揺れはある程度吸収する。
"""

import argparse
import csv
import re
import sys
import unicodedata

# 一括取得サービスごとの列名の揺れを吸収するためのエイリアス
COL_ALIASES = {
    "所在": ["所在", "所在地", "土地所在", "物件所在"],
    "地番": ["地番", "地番号"],
    "所有者": ["所有者", "所有者氏名", "氏名", "名義人", "所有者名"],
    "所有者住所": ["所有者住所", "住所", "所有者の住所", "名義人住所"],
    "地目": ["地目"],
    "地積": ["地積", "公簿地積", "地積㎡"],
}

CORP_PAT = re.compile(r"(株式会社|有限会社|合同会社|合資会社|合名会社|一般社団|一般財団|"
                      r"公益社団|公益財団|医療法人|社会福祉法人|学校法人|宗教法人|農業協同組合|"
                      r"独立行政法人|国土交通省|財務省|法務局|市$|県$|町$|村$)")

TSUBO_PER_M2 = 1 / 3.305785

KANJI_DIGITS = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
                "六": 6, "七": 7, "八": 8, "九": 9}


def _kanji_chome_to_arabic(m):
    """「一丁目」〜「十九丁目」をアラビア数字に(登記とサービスで表記が割れるため)。"""
    s = m.group(1)
    if s == "十":
        n = 10
    elif s.startswith("十"):
        n = 10 + KANJI_DIGITS[s[1]]
    elif s.endswith("十"):
        n = KANJI_DIGITS[s[0]] * 10
    else:
        n = 0
        for ch in s:
            n = n * 10 + KANJI_DIGITS[ch]
    return f"{n}丁目"


def norm(s):
    """所在・地番の照合キー正規化: 全角半角統一・空白除去・丁目の漢数字→算用数字。"""
    if s is None:
        return ""
    s = unicodedata.normalize("NFKC", str(s))
    s = re.sub(r"\s+", "", s)
    return re.sub(r"([一二三四五六七八九十]{1,3})丁目", _kanji_chome_to_arabic, s)


def read_owner_csv(path):
    """所有者CSVを読み、(所在,地番)→レコード の辞書を返す。"""
    raw = open(path, "rb").read()
    text = None
    for enc in ("utf-8-sig", "cp932", "utf-8"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        sys.exit(f"エラー: {path} の文字コードを判定できません(UTF-8/Shift_JISで保存してください)")

    rows = list(csv.reader(text.splitlines()))
    if not rows:
        sys.exit(f"エラー: {path} が空です")

    header = [norm(h) for h in rows[0]]
    colmap = {}
    for key, aliases in COL_ALIASES.items():
        for i, h in enumerate(header):
            if h in aliases:
                colmap[key] = i
                break
    for required in ("所在", "地番", "所有者"):
        if required not in colmap:
            sys.exit(f"エラー: 所有者CSVに「{required}」列が見つかりません。"
                     f"ヘッダー: {rows[0]}")

    owners = {}
    for r in rows[1:]:
        if not any(r):
            continue
        def get(key):
            i = colmap.get(key)
            return r[i].strip() if i is not None and i < len(r) else ""
        key = (norm(get("所在")), norm(get("地番")))
        owners[key] = {
            "所有者": get("所有者"),
            "所有者住所": get("所有者住所"),
            "地目": get("地目"),
            "地積": get("地積"),
        }
    return owners


def main():
    ap = argparse.ArgumentParser(description="所有者CSVをExcelに統合しDMリストを作る")
    ap.add_argument("xlsx", help="extract_parcels.py が出力したExcel")
    ap.add_argument("owners_csv", help="一括取得サービスから返ってきた所有者CSV")
    ap.add_argument("-o", "--output", default="DMリスト.xlsx", help="出力Excel名")
    ap.add_argument("--include-corp", action="store_true",
                    help="法人・官公庁名義もDMリストに含める(既定は個人のみ)")
    args = ap.parse_args()

    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    owners = read_owner_csv(args.owners_csv)
    wb = load_workbook(args.xlsx)
    ws = wb["抽出リスト"] if "抽出リスト" in wb.sheetnames else wb.active

    header = [c.value for c in ws[1]]
    try:
        i_shozai = header.index("所在")
        i_chiban = header.index("地番")
        i_tsubo = header.index("図上坪数")
        i_owner = header.index("所有者(謄本取得後に記入)")
        i_addr = header.index("所有者住所")
        i_chimoku = header.index("地目")
    except ValueError as e:
        sys.exit(f"エラー: Excelの列構成が想定と違います({e})。"
                 "extract_parcels.py の出力ファイルを指定してください。")

    n_hit = n_miss = 0
    by_owner = {}  # (所有者,住所) → {"筆": [...], "坪": float, "法人": bool}
    for row in ws.iter_rows(min_row=2):
        key = (norm(row[i_shozai].value), norm(row[i_chiban].value))
        rec = owners.get(key)
        if not rec:
            n_miss += 1
            continue
        n_hit += 1
        row[i_owner].value = rec["所有者"]
        row[i_addr].value = rec["所有者住所"]
        if rec["地目"]:
            row[i_chimoku].value = rec["地目"]
        okey = (rec["所有者"], rec["所有者住所"])
        d = by_owner.setdefault(okey, {"筆": [], "坪": 0.0,
                                       "法人": bool(CORP_PAT.search(rec["所有者"]))})
        d["筆"].append(f"{row[i_shozai].value} {row[i_chiban].value}")
        d["坪"] += float(row[i_tsubo].value or 0)

    # 所有者別DMリストのシートを作る(同一名義の複数筆は1行にまとめる)
    if "DMリスト" in wb.sheetnames:
        del wb["DMリスト"]
    dm = wb.create_sheet("DMリスト")
    dm.append(["No", "所有者", "所有者住所", "区分", "所有筆数", "合計坪数",
               "所有地(所在 地番)", "DM発送日", "反応", "備考"])
    head_fill = PatternFill("solid", fgColor="1F4E79")
    for c in dm[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
    n_dm = 0
    for (name, addr), d in sorted(by_owner.items(), key=lambda x: -x[1]["坪"]):
        if d["法人"] and not args.include_corp:
            continue
        n_dm += 1
        dm.append([n_dm, name, addr, "法人" if d["法人"] else "個人",
                   len(d["筆"]), round(d["坪"], 1), " / ".join(d["筆"]), "", "", ""])
    for col, w in zip("ABCDEFGHIJ", [5, 22, 40, 6, 8, 10, 60, 12, 10, 20]):
        dm.column_dimensions[col].width = w
    dm.freeze_panes = "A2"
    dm.auto_filter.ref = f"A1:J{n_dm + 1}"

    wb.save(args.output)
    n_corp = sum(1 for d in by_owner.values() if d["法人"])
    print(f"照合: {n_hit}筆ヒット / {n_miss}筆未取得")
    print(f"所有者: {len(by_owner)}名義(うち法人・官公庁 {n_corp})"
          f" → DMリスト {n_dm}行" + ("" if args.include_corp else " (個人のみ)"))
    print(f"出力: {args.output}(「抽出リスト」に所有者を転記+「DMリスト」シートを追加)")


if __name__ == "__main__":
    main()
